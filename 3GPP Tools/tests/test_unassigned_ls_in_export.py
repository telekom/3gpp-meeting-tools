import sys
import unittest
from pathlib import Path
from PyQt5.QtWidgets import QApplication

from modules.meetings.core.post_meeting_exporter import PostMeetingExporterThread
from modules.meetings.core.settings import MeetingsSettings
from modules.meetings.ui.post_meeting_export_dialog import PostMeetingExportDialog
from modules.meetings.ui.tdocs_models import TDocsTableModel, TDocsFilterProxyModel


class TestUnassignedLsInExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not QApplication.instance():
            cls.app = QApplication(sys.argv)
        else:
            cls.app = QApplication.instance()

    def setUp(self):
        self.rows = [
            # 1. LS IN with NO WI assigned and generic/unrelated AI description
            {
                "TDoc": "S2-2609300",
                "Type": "LS in",
                "TDoc Status": "",
                "Related WIs": "",
                "Agenda Item": "19.2.2",
                "Agenda item description": "System architecture for next generation real time communication (NG_RTC_Ph2)",
                "Release": "Rel-19",
                "Title": "LS on NG-RTC",
            },
            # 2. LS IN with NO WI assigned and AI description containing parenthetical CR(s)
            {
                "TDoc": "S2-2609338",
                "Type": "LS in",
                "TDoc Status": "",
                "Related WIs": "",
                "Agenda Item": "20.50",
                "Agenda item description": "Rel-20 LSs and CAT B/C alignment CR(s) due to the study",
                "Release": "Rel-20",
                "Title": "LS on alignment",
            },
            # 3. LS IN with NO WI assigned and common issues AI 4.1
            {
                "TDoc": "S2-2609305",
                "Type": "LS in",
                "TDoc Status": "",
                "Related WIs": "",
                "Agenda Item": "4.1",
                "Agenda item description": "Common issues and Incoming LSs",
                "Release": "Rel-20",
                "Title": "General Incoming LS",
            },
            # 4. LS IN with NO WI assigned but AI description matches followed WI
            {
                "TDoc": "S2-2609303",
                "Type": "LS in",
                "TDoc Status": "",
                "Related WIs": "",
                "Agenda Item": "20.6.20",
                "Agenda item description": "Study on Architecture for 6G System (FS_6G_ARC)",
                "Release": "Rel-20",
                "Title": "LS on 6G Arch",
            },
            # 5. LS IN with an ASSIGNED non-followed WI
            {
                "TDoc": "S2-2609399",
                "Type": "LS in",
                "TDoc Status": "",
                "Related WIs": "AmbientIoT_Ph2",
                "Agenda Item": "20.5.2",
                "Agenda item description": "Ambient IoT",
                "Release": "Rel-20",
                "Title": "Assigned LS to Ambient IoT",
            },
            # 6. Followed WI CR (Agreed)
            {
                "TDoc": "S2-2609400",
                "Type": "CR",
                "TDoc Status": "Agreed",
                "Related WIs": "6G_ARC",
                "Agenda Item": "20.6.20",
                "Agenda item description": "Study on Architecture for 6G System (FS_6G_ARC)",
                "Release": "Rel-20",
                "Title": "Agreed CR on 6G",
            },
            # 7. Non-followed WI CR (Agreed)
            {
                "TDoc": "S2-2609401",
                "Type": "CR",
                "TDoc Status": "Agreed",
                "Related WIs": "AmbientIoT_Ph2",
                "Agenda Item": "20.5.2",
                "Agenda item description": "Ambient IoT",
                "Release": "Rel-20",
                "Title": "Agreed CR on Ambient IoT",
            },
        ]
        # User only follows 6G_ARC
        self.followed_tokens = {"6g_arc", "fs_6g_arc"}
        self.wi_db_info = {
            "wi_by_exact_acronym": {
                "fs_6g_arc": {"acronym": "FS_6G_ARC", "code": "1080058", "name": "6G Architecture", "release": "Rel-20"},
                "6g_arc": {"acronym": "6G_ARC", "code": "1080058", "name": "6G Architecture", "release": "Rel-20"},
                "ng_rtc_ph2": {"acronym": "NG_RTC_Ph2", "code": "9999", "name": "NG RTC", "release": "Rel-19"},
                "ambientiot_ph2": {"acronym": "AmbientIoT_Ph2", "code": "8888", "name": "Ambient IoT", "release": "Rel-20"},
            },
            "wi_by_code": {},
            "wi_by_token": {},
            "followed_tokens": self.followed_tokens,
            "starred_tokens": set(),
            "tracked_tokens": self.followed_tokens,
        }

    def test_post_meeting_exporter_includes_unassigned_ls_in_by_default(self):
        thread = PostMeetingExporterThread(
            meeting_dir=Path("."),
            all_rows=self.rows,
            visible_rows=self.rows,
            agenda_map={},
            mtg_info={"wg_name": "SA2", "meeting_number": "177"},
            docs_ftp_url="",
            revisions_url="",
            scope_mode="followed",
            followed_tokens=self.followed_tokens,
            include_ls_in=True,
            include_general_ls=True,
        )

        tree = thread._filter_and_group_tdocs(self.wi_db_info)

        # Collect all exported TDoc IDs by category and WI
        exported_tdocs = {}
        for rel, wis in tree.items():
            for wi, data in wis.items():
                for cat, doc_list in data["categories"].items():
                    for d in doc_list:
                        exported_tdocs[d["tdoc_id"]] = (rel, wi, cat)

        # 1. S2-2609300 (LS IN, no WI in Related WIs, but NG_RTC_Ph2 in AIDesc)
        # MUST be included under General/LS_IN because it has no official WI and NG_RTC_Ph2 is not followed!
        self.assertIn("S2-2609300", exported_tdocs)
        self.assertEqual(exported_tdocs["S2-2609300"][1], "General")
        self.assertEqual(exported_tdocs["S2-2609300"][2], "LS_IN")

        # 2. S2-2609338 (LS IN, no WI, CR(s) in AIDesc)
        # MUST be included under General/LS_IN!
        self.assertIn("S2-2609338", exported_tdocs)
        self.assertEqual(exported_tdocs["S2-2609338"][1], "General")
        self.assertEqual(exported_tdocs["S2-2609338"][2], "LS_IN")

        # 3. S2-2609305 (AI 4.1 unassigned LS IN)
        # MUST be included under General/LS_IN!
        self.assertIn("S2-2609305", exported_tdocs)
        self.assertEqual(exported_tdocs["S2-2609305"][1], "General")

        # 4. S2-2609303 (LS IN matching followed WI 6G_ARC)
        # MUST be included under FS_6G_ARC!
        self.assertIn("S2-2609303", exported_tdocs)
        self.assertEqual(exported_tdocs["S2-2609303"][1], "FS_6G_ARC")

        # 5. S2-2609399 (LS IN with an ASSIGNED non-followed WI)
        # Should be excluded because it has a specific assigned WI not followed by the user
        self.assertNotIn("S2-2609399", exported_tdocs)

        # 6. S2-2609400 (Agreed CR on followed WI)
        # Included under 6G_ARC
        self.assertIn("S2-2609400", exported_tdocs)
        self.assertEqual(exported_tdocs["S2-2609400"][1], "6G_ARC")

        # 7. S2-2609401 (Agreed CR on non-followed WI)
        # Excluded
        self.assertNotIn("S2-2609401", exported_tdocs)

    def test_post_meeting_exporter_excludes_unassigned_ls_in_when_disabled(self):
        thread = PostMeetingExporterThread(
            meeting_dir=Path("."),
            all_rows=self.rows,
            visible_rows=self.rows,
            agenda_map={},
            mtg_info={"wg_name": "SA2", "meeting_number": "177"},
            docs_ftp_url="",
            revisions_url="",
            scope_mode="followed",
            followed_tokens=self.followed_tokens,
            include_ls_in=True,
            include_general_ls=False,  # Disabled
        )

        tree = thread._filter_and_group_tdocs(self.wi_db_info)
        exported_ids = set()
        for rel, wis in tree.items():
            for wi, data in wis.items():
                for cat, doc_list in data["categories"].items():
                    for d in doc_list:
                        exported_ids.add(d["tdoc_id"])

        # Unassigned LS INs not matching followed WIs must now be excluded
        self.assertNotIn("S2-2609300", exported_ids)
        self.assertNotIn("S2-2609338", exported_ids)
        self.assertNotIn("S2-2609305", exported_ids)

        # But LS IN matching followed WI must still be included
        self.assertIn("S2-2609303", exported_ids)
        self.assertIn("S2-2609400", exported_ids)

    def test_post_meeting_export_dialog_defaults(self):
        dialog = PostMeetingExportDialog(
            parent=None,
            meeting_dir=Path("."),
            all_rows=self.rows,
            visible_rows=self.rows,
            agenda_map={},
            mtg_info={"wg_name": "SA2", "meeting_number": "177"},
            followed_tokens=self.followed_tokens,
            starred_tokens=set(),
        )
        cfg = dialog.get_export_config()
        self.assertTrue(cfg["include_ls_in"])
        self.assertTrue(cfg["include_general_ls"])
        self.assertEqual(cfg["scope_mode"], "followed")

    def test_proxy_model_tracked_wi_filtering_with_unassigned_ls_in(self):
        model = TDocsTableModel(Path("."), self.rows)
        proxy = TDocsFilterProxyModel()
        proxy.setSourceModel(model)

        # Initialize filter sets
        proxy.setTypeFilters({str(r.get("Type", "")).strip() for r in self.rows})
        proxy.setAIFilters({str(r.get("Agenda Item", "")).strip() for r in self.rows})
        proxy.setWIFilters(set().union(*[model.get_related_wis(i) for i in range(len(self.rows))]))
        proxy.setStatusFilters({str(r.get("TDoc Status", "")).strip() for r in self.rows})
        proxy.setMyStatusFilters({str(r.get("My Status", "")).strip() for r in self.rows})
        proxy.setCompanyFilters(set().union(*[r.get("_Sanitized_Companies", ["Other"]) for r in self.rows]))

        proxy.setTrackedWITokens(self.followed_tokens, set())
        proxy.setTrackedWIFilters(True, False)

        # Default: include_unassigned_ls_in is True
        self.assertTrue(proxy.include_unassigned_ls_in)

        visible_ids = set()
        for r in range(proxy.rowCount()):
            src_idx = proxy.mapToSource(proxy.index(r, 0))
            visible_ids.add(model._data[src_idx.row()]["TDoc"])

        # Unassigned LS INs must be visible
        self.assertIn("S2-2609300", visible_ids)
        self.assertIn("S2-2609338", visible_ids)
        self.assertIn("S2-2609305", visible_ids)
        self.assertIn("S2-2609303", visible_ids)
        # Followed CR must be visible
        self.assertIn("S2-2609400", visible_ids)

        # Non-followed assigned items must NOT be visible
        self.assertNotIn("S2-2609399", visible_ids)
        self.assertNotIn("S2-2609401", visible_ids)

        # When toggled off:
        proxy.setIncludeUnassignedLsIn(False)
        visible_ids_off = set()
        for r in range(proxy.rowCount()):
            src_idx = proxy.mapToSource(proxy.index(r, 0))
            visible_ids_off.add(model._data[src_idx.row()]["TDoc"])

        self.assertNotIn("S2-2609300", visible_ids_off)
        self.assertNotIn("S2-2609338", visible_ids_off)
        self.assertNotIn("S2-2609305", visible_ids_off)
        self.assertIn("S2-2609303", visible_ids_off)
        self.assertIn("S2-2609400", visible_ids_off)

    def test_settings_include_unassigned_ls_in_persistence(self):
        settings = MeetingsSettings()
        initial = settings.get_include_unassigned_ls_in()
        self.assertIsInstance(initial, bool)

        # Toggle and verify persistence
        settings.set_include_unassigned_ls_in(False)
        self.assertFalse(settings.get_include_unassigned_ls_in())

        settings.set_include_unassigned_ls_in(True)
        self.assertTrue(settings.get_include_unassigned_ls_in())

    def test_post_meeting_exporter_plain_document_handling(self):
        """Verify that PostMeetingExporter renders LS IN documents in accept mode without change marks."""
        import tempfile, shutil
        from modules.meetings.core.tdoc_markdown_converter import resolve_revision_mode_for_type

        tmp_dir = Path(tempfile.mkdtemp(prefix="test_pm_plain_"))
        try:
            meeting_dir = tmp_dir / "Meeting"
            meeting_dir.mkdir(parents=True, exist_ok=True)

            exporter = PostMeetingExporterThread(
                meeting_dir=meeting_dir,
                all_rows=[],
                visible_rows=[],
                agenda_map={},
                mtg_info={},
                docs_ftp_url="",
                revisions_url="",
                scope_mode="whole",
                output_dir=tmp_dir / "Export",
            )

            # 1. Mode resolution check
            self.assertEqual(resolve_revision_mode_for_type("LS in"), "accept")
            self.assertEqual(resolve_revision_mode_for_type("LS out"), "accept")
            self.assertEqual(resolve_revision_mode_for_type("CR"), "markup")
            self.assertEqual(resolve_revision_mode_for_type("pCR"), "markup")

            # 2. Test docx conversion via _ensure_tdoc_markdown
            import docx
            tdoc_folder = meeting_dir / "S2-2609001"
            tdoc_folder.mkdir(parents=True, exist_ok=True)
            doc = docx.Document()
            doc.add_heading("Incoming Liaison", level=1)
            doc.add_paragraph("This is a clean communication from another WG.")
            docx_path = tdoc_folder / "S2-2609001.docx"
            doc.save(str(docx_path))

            md_path, _ = exporter._ensure_tdoc_markdown(
                "S2-2609001",
                {"TDoc": "S2-2609001", "Type": "LS in", "Title": "Incoming Liaison"}
            )
            self.assertTrue(md_path.exists())
            text = md_path.read_text(encoding="utf-8")
            self.assertIn("mode=accept", text)
            self.assertIn("(revisions: accept)", text)
            self.assertNotIn("Markup legend:", text)
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    def test_post_meeting_ls_reply_tracking_and_inclusion(self):
        import tempfile
        import shutil

        rows = [
            # 1. Incoming LS with no assigned WI, replied to via S2-2610500 (draft)
            {
                "TDoc": "S2-2609300",
                "Type": "LS in",
                "TDoc Status": "replied to",
                "Reply in": "S2-2610500, S2-2610500",
                "Reply to": "",
                "Related WIs": "",
                "Agenda Item": "4.1",
                "Agenda item description": "Common issues and Incoming LSs",
                "Release": "Rel-19",
                "Title": "Incoming LS on Avatars",
                "Source": "GSMA",
            },
            # 2. Draft reply LS, revised to S2-2610956
            {
                "TDoc": "S2-2610500",
                "Type": "LS out",
                "TDoc Status": "revised",
                "Revised to": "S2-2610956",
                "Reply in": "",
                "Reply to": "S2-2609300",
                "Related WIs": "AvCall-MED",
                "Agenda Item": "19.2",
                "Agenda item description": "Avatars Call",
                "Release": "Rel-19",
                "Title": "Draft Reply LS",
                "Source": "SA2",
            },
            # 3. Final approved reply LS (WI is AvCall-MED, NOT in user followed tokens)
            {
                "TDoc": "S2-2610956",
                "Type": "LS out",
                "TDoc Status": "approved",
                "Reply in": "",
                "Reply to": "S2-2609300",
                "Related WIs": "AvCall-MED",
                "Agenda Item": "19.2",
                "Agenda item description": "Avatars Call",
                "Release": "Rel-19",
                "Title": "Final Approved Reply LS",
                "Source": "SA2",
            },
            # 4. Followed WI CR (User only follows 6G_ARC)
            {
                "TDoc": "S2-2609400",
                "Type": "CR",
                "TDoc Status": "Agreed",
                "Related WIs": "6G_ARC",
                "Agenda Item": "20.6.20",
                "Agenda item description": "Study on 6G",
                "Release": "Rel-20",
                "Title": "Agreed CR on 6G",
                "Source": "Qualcomm",
            },
        ]

        tmp_dir = Path(tempfile.mkdtemp(prefix="test_ls_replies_"))
        try:
            meeting_dir = tmp_dir / "Meeting"
            export_dir = tmp_dir / "Export"
            meeting_dir.mkdir(parents=True, exist_ok=True)
            export_dir.mkdir(parents=True, exist_ok=True)

            exporter = PostMeetingExporterThread(
                meeting_dir=meeting_dir,
                all_rows=rows,
                visible_rows=rows,
                agenda_map={},
                mtg_info={"wg_name": "SA2", "meeting_number": "177"},
                docs_ftp_url="",
                revisions_url="",
                scope_mode="followed",
                followed_tokens={"6g_arc", "fs_6g_arc"},
                include_ls_in=True,
                include_ls_out=True,
                include_general_ls=True,
                output_dir=export_dir,
            )

            # 1. Test reply mapping with revision resolution
            in_to_replies, out_to_in, tdoc_map = exporter._build_ls_reply_map()
            self.assertIn("S2-2609300", in_to_replies)
            self.assertEqual(in_to_replies["S2-2609300"], {"S2-2610956"})
            self.assertIn("S2-2610956", out_to_in)
            self.assertEqual(out_to_in["S2-2610956"], {"S2-2609300"})

            # 2. Test grouping and guaranteed inclusion of replies
            tree = exporter._filter_and_group_tdocs(self.wi_db_info)

            # S2-2609300 (LS IN) must be in Rel-19/General
            self.assertIn("Rel-19", tree)
            self.assertIn("General", tree["Rel-19"])
            ls_in_docs = tree["Rel-19"]["General"]["categories"]["LS_IN"]
            self.assertEqual(len(ls_in_docs), 1)
            self.assertEqual(ls_in_docs[0]["tdoc_id"], "S2-2609300")
            self.assertEqual(ls_in_docs[0]["replies"], ["S2-2610956"])

            # S2-2610956 (LS OUT) MUST be in Rel-19/General/LS_OUT even though AvCall-MED is not followed!
            ls_out_docs = tree["Rel-19"]["General"]["categories"]["LS_OUT"]
            self.assertEqual(len(ls_out_docs), 1)
            self.assertEqual(ls_out_docs[0]["tdoc_id"], "S2-2610956")
            self.assertEqual(ls_out_docs[0]["reply_to_ids"], ["S2-2609300"])
            self.assertEqual(ls_out_docs[0]["cross_wis"], ["AvCall-MED"])

            # AvCall-MED must NOT be present as an independent WI section
            self.assertNotIn("AvCall-MED", tree.get("Rel-19", {}))
            self.assertNotIn("AvCall-MED", tree.get("Rel-20", {}))

            # Followed WI 6G_ARC must be present in Rel-20
            self.assertIn("Rel-20", tree)
            self.assertIn("6G_ARC", tree["Rel-20"])

            # 3. Test dossier export and reciprocal headers
            # Create dummy cache files
            for t_id, title in [("S2-2609300", "Incoming LS on Avatars"), ("S2-2610956", "Final Approved Reply LS")]:
                f_dir = meeting_dir / t_id
                f_dir.mkdir(parents=True, exist_ok=True)
                (f_dir / f"{t_id}.md").write_text(f"# {t_id}: {title}\n\nDocument body.", encoding="utf-8")

            exporter._export_all_dossiers(tree)

            in_md_file = export_dir / "Rel-19" / "General" / "LS_IN" / "S2-2609300.md"
            out_md_file = export_dir / "Rel-19" / "General" / "LS_OUT" / "S2-2610956.md"
            self.assertTrue(in_md_file.exists())
            self.assertTrue(out_md_file.exists())

            in_content = in_md_file.read_text(encoding="utf-8")
            out_content = out_md_file.read_text(encoding="utf-8")

            # Check reciprocal callouts
            self.assertIn("[!TIP]", in_content)
            self.assertIn("Outgoing Reply:", in_content)
            self.assertIn("[S2-2610956](../LS_OUT/S2-2610956.md)", in_content)

            self.assertIn("[!NOTE]", out_content)
            self.assertIn("Reply Liaison Statement:", out_content)
            self.assertIn("[S2-2609300](../LS_IN/S2-2609300.md)", out_content)

            # 4. Test table rendering with reply hyperlinks
            # Master index table for LS_IN
            master_in_table = exporter._render_markdown_table(ls_in_docs, is_local_wi=False, base_dir=export_dir)
            self.assertIn("replied to in [S2-2610956](Rel-19/General/LS_OUT/S2-2610956.md)", master_in_table)

            # Master index table for LS_OUT
            master_out_table = exporter._render_markdown_table(ls_out_docs, is_local_wi=False, base_dir=export_dir)
            self.assertIn("approved (reply to [S2-2609300](Rel-19/General/LS_IN/S2-2609300.md))", master_out_table)

            # Local WI index table for LS_IN (from Rel-19/General/index.md)
            general_dir = export_dir / "Rel-19" / "General"
            local_in_table = exporter._render_markdown_table(ls_in_docs, is_local_wi=True, base_dir=general_dir)
            self.assertIn("replied to in [S2-2610956](LS_OUT/S2-2610956.md)", local_in_table)

            # Local WI index table for LS_OUT
            local_out_table = exporter._render_markdown_table(ls_out_docs, is_local_wi=True, base_dir=general_dir)
            self.assertIn("approved (reply to [S2-2609300](LS_IN/S2-2609300.md))", local_out_table)

        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    def test_post_meeting_exporter_parallel_predownload(self):
        """Verify PostMeetingExporterThread parallel pre-download triggers for missing TDocs."""
        import shutil
        import tempfile
        from unittest.mock import patch

        tmp_dir = Path(tempfile.mkdtemp(prefix="3gpp_test_pm_dl_"))
        try:
            meeting_dir = tmp_dir / "Meeting_102"
            meeting_dir.mkdir(parents=True, exist_ok=True)

            rows = [
                {"TDoc": "S2-2601111", "Type": "CR", "TDoc Status": "agreed", "Related WIs": "WI_TEST", "Release": "Rel-19", "Title": "CR 1"},
                {"TDoc": "S2-2602222", "Type": "CR", "TDoc Status": "agreed", "Related WIs": "WI_TEST", "Release": "Rel-19", "Title": "CR 2"},
            ]

            exporter = PostMeetingExporterThread(
                meeting_dir=meeting_dir,
                all_rows=rows,
                visible_rows=rows,
                agenda_map={},
                mtg_info={"wg_name": "SA2", "meeting_number": "177"},
                docs_ftp_url="https://www.3gpp.org/ftp/tsg_sa/WG2_Arch/Docs",
                revisions_url="https://www.3gpp.org/ftp/tsg_sa/WG2_Arch/INBOX/Revisions",
                scope_mode="all",
                download_missing=True,
            )

            downloaded = []
            def fake_download(tdoc_id, folder):
                downloaded.append(tdoc_id)
                folder.mkdir(parents=True, exist_ok=True)
                doc_path = folder / f"{tdoc_id}.docx"
                doc_path.write_bytes(b"dummy")
                md_path = folder / f"{tdoc_id}.md"
                md_path.write_text(f"# {tdoc_id}\nContent", encoding="utf-8")
                return doc_path

            progress_messages = []
            exporter.progress.connect(progress_messages.append)

            with patch.object(exporter, "_download_and_extract_tdoc", side_effect=fake_download):
                exporter.run()

            self.assertIn("S2-2601111", downloaded)
            self.assertIn("S2-2602222", downloaded)
            joined_prog = " ".join(progress_messages)
            self.assertIn("in parallel", joined_prog)

        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
