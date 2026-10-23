import os
import sys
import tempfile
import time
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import openpyxl
from PyQt5.QtCore import QCoreApplication

from modules.meetings.core.meeting_wi_enricher import (
    MeetingWIEnricherThread,
    extract_wi_ids_from_excel,
    get_marker_path,
    is_enrichment_needed,
    write_enrichment_marker,
)
from modules.meetings.core.settings import MeetingsSettings
from modules.work_items.core.wi_database import WorkItemsDatabase


class TestMeetingWIEnricher(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not QCoreApplication.instance():
            cls.app = QCoreApplication(sys.argv)
        else:
            cls.app = QCoreApplication.instance()

    def test_marker_path_and_enrichment_needed(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            excel_path = Path(tmp_dir) / "TDoc_List_Test.xlsx"
            excel_path.write_text("dummy excel content")

            marker_path = get_marker_path(excel_path)
            self.assertEqual(marker_path, Path(tmp_dir) / "TDoc_List_Test.xlsx.wi_checked.json")

            # Initially needed because marker does not exist
            self.assertTrue(is_enrichment_needed(excel_path))

            # Write marker
            write_enrichment_marker(excel_path, unique_wis_found=5, new_wis_scraped=2)
            self.assertTrue(marker_path.exists())

            # Now not needed
            self.assertFalse(is_enrichment_needed(excel_path))

            # Touch excel to make it newer than marker
            time.sleep(0.05)
            excel_path.write_text("updated excel content")
            self.assertTrue(is_enrichment_needed(excel_path))

    def test_extract_wi_ids_from_excel(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            excel_path = Path(tmp_dir) / "test_meeting.xlsx"

            # Create an in-memory openpyxl workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "TDoc_List"

            # Row 1: Headers
            ws.cell(row=1, column=1, value="AI")
            ws.cell(row=1, column=2, value="TDoc")
            ws.cell(row=1, column=3, value="Title")
            ws.cell(row=1, column=4, value="Related WIs")

            # Row 2: Single WI with hyperlink
            ws.cell(row=2, column=2, value="S2-260001")
            c1 = ws.cell(row=2, column=4, value="5GSAT_ARCH")
            c1.hyperlink = "https://portal.3gpp.org/desktopmodules/WorkItem/WorkItemDetails.aspx?workitemId=890034"

            # Row 3: Single WI with hyperlink
            ws.cell(row=3, column=2, value="S2-260002")
            c2 = ws.cell(row=3, column=4, value="AIML_CN_Ph2-CT")
            c2.hyperlink = "https://portal.3gpp.org/desktopmodules/WorkItem/WorkItemDetails.aspx?workitemId=1120028"

            # Row 4: Multiple WIs without hyperlink
            ws.cell(row=4, column=2, value="S2-260003")
            ws.cell(row=4, column=4, value="TEI19, TEI20")

            # Row 5: Empty TDoc row to trigger break
            wb.save(excel_path)
            wb.close()

            found = extract_wi_ids_from_excel(excel_path)
            self.assertEqual(len(found), 2)
            self.assertIn("890034", found)
            self.assertIn("1120028", found)
            self.assertEqual(found["890034"], "5GSAT_ARCH")
            self.assertEqual(found["1120028"], "AIML_CN_Ph2-CT")

    def test_database_get_existing_wi_codes(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            db_path = Path(tmp_dir) / "test_3gpp.db"
            db = WorkItemsDatabase(db_path)

            codes_to_check = {"890034", "9999999", "1120028"}
            existing = db.get_existing_wi_codes(codes_to_check)
            self.assertEqual(existing, set())

            # Upsert a single WI
            db.upsert_single_work_item({
                "code": "890034",
                "acronym": "5GSAT_ARCH",
                "name": "Satellite Architecture",
                "release": "Rel-18",
            })

            existing = db.get_existing_wi_codes(codes_to_check)
            self.assertEqual(existing, {"890034"})

    def test_enricher_thread_when_already_up_to_date(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            excel_path = Path(tmp_dir) / "test_up_to_date.xlsx"
            excel_path.write_text("sample content")

            write_enrichment_marker(excel_path, unique_wis_found=0, new_wis_scraped=0)

            thread = MeetingWIEnricherThread(excel_path, db_path=Path(tmp_dir) / "empty.db")
            results = []
            thread.finished_enrichment.connect(lambda s, c, m: results.append((s, c, m)))
            thread.start()
            thread.wait(2000)
            QCoreApplication.processEvents()

            self.assertEqual(len(results), 1)
            self.assertTrue(results[0][0])
            self.assertEqual(results[0][1], 0)
            self.assertIn("already", results[0][2].lower())


if __name__ == "__main__":
    unittest.main()
