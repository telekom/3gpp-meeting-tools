import tempfile
import unittest
from pathlib import Path

import docx

from modules.meetings.core.tdocs_parser import TDocsParser


class TestChairNotesWordParser(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_dir = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_parse_sa2_standard_format(self):
        """Tests standard SA2 Chairman's notes / TdocsByAgenda table structure."""
        doc = docx.Document()
        table = doc.add_table(rows=3, cols=8)

        # Header on row 0
        headers = ['TD#', 'AI', 'Title', 'Source', 'Type', 'For', 'Comments', 'Result']
        for i, h in enumerate(headers):
            table.cell(0, i).text = h

        # Row 1
        r1 = ['S2-2401234', '5.1', 'Architecture enhancements for IMS', 'Deutsche Telekom', 'CR', 'Agreement', 'Discussion in room: minor edits requested', 'Revised to S2-2401890']
        for i, val in enumerate(r1):
            table.cell(1, i).text = val

        # Row 2
        r2 = ['S2-2401235', '5.1', 'Discussion on QoS flows', 'Qualcomm', 'Discussion', 'Discussion', 'Agreed offline', 'Agreed']
        for i, val in enumerate(r2):
            table.cell(2, i).text = val

        file_path = self.test_dir / "sa2_notes.docx"
        doc.save(str(file_path))

        parsed = TDocsParser.parse_tdocs_from_docx(str(file_path))
        self.assertEqual(len(parsed), 2)
        self.assertIn('S2-2401234', parsed)
        self.assertEqual(parsed['S2-2401234']['Agenda Item'], '5.1')
        self.assertEqual(parsed['S2-2401234']['Title'], 'Architecture enhancements for IMS')
        self.assertEqual(parsed['S2-2401234']['Source'], 'Deutsche Telekom')
        self.assertEqual(parsed['S2-2401234']['Comments'], 'Discussion in room: minor edits requested')
        self.assertEqual(parsed['S2-2401234']['Result'], 'Revised to S2-2401890')

    def test_parse_ran_format_with_banner_and_discussion_column(self):
        """Tests RAN2/RAN3 format where row 0 is a meeting banner and column is 'Discussion'."""
        doc = docx.Document()
        table = doc.add_table(rows=3, cols=6)

        # Row 0: Meeting title banner
        for c in range(6):
            table.cell(0, c).text = "3GPP TSG-RAN WG2 Meeting #120 - Chairman's Notes"

        # Row 1: True headers
        headers = ['TDoc #', 'Agenda Item', 'Document Title', 'Company', 'Discussion', 'Decision']
        for i, h in enumerate(headers):
            table.cell(1, i).text = h

        # Row 2: Data row
        r2 = ['R2-2405678', '6.2.1', 'NR Sidelink Relay Configurations', 'Ericsson', 'No objections heard. Endorsed by room.', 'Endorsed']
        for i, val in enumerate(r2):
            table.cell(2, i).text = val

        file_path = self.test_dir / "ran_notes.docx"
        doc.save(str(file_path))

        parsed = TDocsParser.parse_tdocs_from_docx(str(file_path))
        self.assertEqual(len(parsed), 1)
        self.assertIn('R2-2405678', parsed)
        self.assertEqual(parsed['R2-2405678']['Agenda Item'], '6.2.1')
        self.assertEqual(parsed['R2-2405678']['Title'], 'NR Sidelink Relay Configurations')
        self.assertEqual(parsed['R2-2405678']['Source'], 'Ericsson')
        self.assertEqual(parsed['R2-2405678']['Comments'], 'No objections heard. Endorsed by room.')
        self.assertEqual(parsed['R2-2405678']['Result'], 'Endorsed')

    def test_parse_section_divider_and_tdoc_token_extraction(self):
        """Tests table with merged section divider rows and decorated TDoc IDs."""
        doc = docx.Document()
        table = doc.add_table(rows=4, cols=5)

        # Headers
        headers = ['Doc #', 'Subject', 'Proposer', 'Notes', 'Status']
        for i, h in enumerate(headers):
            table.cell(0, i).text = h

        # Row 1: Section divider
        for c in range(5):
            table.cell(1, c).text = "Agenda Item 8.3: Positioning & RedCap"

        # Row 2: TDoc with brackets and comment
        r2 = ['[C1-2409999]', 'RedCap UE Capabilities update', 'Nokia', 'Company noted support; offline drafting', 'Postponed']
        for i, val in enumerate(r2):
            table.cell(2, i).text = val

        # Row 3: Revision TDoc with parenthetical note
        r3 = ['C1-2409999r01 (rev of C1-2409999)', 'RedCap UE Capabilities update', 'Nokia', 'Agreed revision', 'Agreed']
        for i, val in enumerate(r3):
            table.cell(3, i).text = val

        file_path = self.test_dir / "ct_notes.docx"
        doc.save(str(file_path))

        parsed = TDocsParser.parse_tdocs_from_docx(str(file_path))
        self.assertEqual(len(parsed), 2)
        self.assertIn('C1-2409999', parsed)
        self.assertEqual(parsed['C1-2409999']['Agenda Item'], '8.3')
        self.assertEqual(parsed['C1-2409999']['Comments'], 'Company noted support; offline drafting')

        self.assertIn('C1-2409999R01', parsed)
        self.assertEqual(parsed['C1-2409999R01']['Agenda Item'], '8.3')
        self.assertEqual(parsed['C1-2409999R01']['Result'], 'Agreed')

    def test_merge_sequential_session_notes(self):
        """Tests that multiple occurrences of a TDoc combine comments."""
        doc = docx.Document()
        table = doc.add_table(rows=3, cols=4)

        headers = ['TD#', 'Subject', 'Discussion', 'Status']
        for i, h in enumerate(headers):
            table.cell(0, i).text = h

        # Morning session
        r1 = ['S2-2408888', 'Network Slice Admission Control', 'Morning: Objections from China Mobile', 'Not Handled']
        for i, val in enumerate(r1):
            table.cell(1, i).text = val

        # Afternoon session
        r2 = ['S2-2408888', 'Network Slice Admission Control', 'Afternoon: Offline consensus reached', 'Agreed']
        for i, val in enumerate(r2):
            table.cell(2, i).text = val

        file_path = self.test_dir / "multi_session.docx"
        doc.save(str(file_path))

        parsed = TDocsParser.parse_tdocs_from_docx(str(file_path))
        self.assertEqual(len(parsed), 1)
        self.assertIn('S2-2408888', parsed)
        self.assertIn('Morning: Objections from China Mobile', parsed['S2-2408888']['Comments'])
        self.assertIn('Afternoon: Offline consensus reached', parsed['S2-2408888']['Comments'])
        self.assertEqual(parsed['S2-2408888']['Result'], 'Agreed')

    def test_parse_ran_tab_delimited_paragraphs(self):
        """Tests paragraph-based format used by RAN1 and RAN2 (no Word tables)."""
        doc = docx.Document()

        doc.add_heading("2.1 Approval of the agenda", level=2)
        doc.add_paragraph("R2-2604501\tAgenda for RAN2#135\tChairman\tagenda")
        p_res = doc.add_paragraph("Approved")
        p_res.style = "Normal"

        doc.add_paragraph("R2-2604502\tRAN2#134 Meeting Report\tMCC\treport")
        doc.add_paragraph("Approved")

        doc.add_heading("4.2 NR Sidelink Enhancements", level=2)
        doc.add_paragraph("R2-2605890\tSidelink relay discovery procedure\tQualcomm\tDiscussion")
        doc.add_paragraph("Room discussion: companies supported Option 2.")
        doc.add_paragraph("Revised to R2-2606500")

        file_path = self.test_dir / "ran_paragraphs.docx"
        doc.save(str(file_path))

        parsed = TDocsParser.parse_tdocs_from_docx(str(file_path))
        self.assertEqual(len(parsed), 3)

        self.assertIn("R2-2604501", parsed)
        self.assertEqual(parsed["R2-2604501"]["Agenda Item"], "2.1")
        self.assertEqual(parsed["R2-2604501"]["Title"], "Agenda for RAN2#135")
        self.assertEqual(parsed["R2-2604501"]["Source"], "Chairman")
        self.assertEqual(parsed["R2-2604501"]["Result"], "Approved")

        self.assertIn("R2-2605890", parsed)
        self.assertEqual(parsed["R2-2605890"]["Agenda Item"], "4.2")
        self.assertEqual(parsed["R2-2605890"]["Title"], "Sidelink relay discovery procedure")
        self.assertIn("Room discussion: companies supported Option 2.", parsed["R2-2605890"]["Comments"])
        self.assertIn("Revised to R2-2606500", parsed["R2-2605890"]["Result"])


if __name__ == '__main__':
    unittest.main()
