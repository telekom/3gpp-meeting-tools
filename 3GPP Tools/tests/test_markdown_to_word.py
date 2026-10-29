"""
Tests for Markdown to Word (.docx) conversion feature.
Verifies Markdown parsing, OpenXML synthesis, media/image embedding,
PlantUML diagram compilation, 3GPP styles, and Studio UI integration.
"""

import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from PyQt5.QtCore import QEventLoop
from PyQt5.QtWidgets import QApplication

# Ensure QApplication exists for UI tests
app = QApplication.instance()
if not app:
    app = QApplication([])

import docx
from core.queue_manager import _TASK_REGISTRY
from modules.word_tools.core.markdown_to_docx import (
    MarkdownToWordThread,
    convert_markdown_to_docx,
    render_plantuml_to_png,
    render_svg_to_png,
    resolve_image_path,
)
from modules.word_tools.plugin_loader import register_word_plugin
from modules.word_tools.ui.cards.markdown_card import WordMarkdownCard
from modules.word_tools.ui.word_tabs import WordExtractorTab


class TestMarkdownToWord(unittest.TestCase):
    """Test suite for Markdown to Word (.docx) feature."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp(prefix="test_md_to_word_")
        self.td = Path(self.temp_dir).resolve()
        register_word_plugin()

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_queue_manager_task_registration(self):
        """Verify that markdown_to_word is registered in the task registry."""
        self.assertIn("markdown_to_word", _TASK_REGISTRY)
        entry = _TASK_REGISTRY["markdown_to_word"]
        self.assertEqual(entry["display_name"], "MARKDOWN → WORD")

        # Verify thread factory produces MarkdownToWordThread
        dummy_file = self.td / "sample.md"
        dummy_file.write_text("# Test", encoding="utf-8")
        thread = entry["factory"](dummy_file, {}, {})
        self.assertIsInstance(thread, MarkdownToWordThread)

    def test_basic_markdown_to_docx_synthesis(self):
        """Verify headings, paragraphs, lists, code, and tables compile to Word."""
        md_file = self.td / "document.md"
        content = """# Heading 1: Overview

This is standard paragraph text with **bold**, *italic*, `monospace code`, and [links](https://3gpp.org).

## Heading 2: Specifications Table

| Spec | Version | Status |
| :--- | :--- | :--- |
| **TS 23.501** | v18.4.0 | Approved |
| **TS 24.501** | v18.4.0 | Approved |

### Heading 3: Requirements List

- Clause requirement A
- Clause requirement B
  - Sub-requirement 1
  - Sub-requirement 2

1. First procedural step
2. Second procedural step

```python
def example_function():
    return "Hello 3GPP"
```
"""
        md_file.write_text(content, encoding="utf-8")
        out_docx = self.td / "document.docx"

        res = convert_markdown_to_docx(
            md_path=md_file,
            output_docx_path=out_docx,
            render_diagrams=False,
            apply_3gpp_styles=True,
        )

        self.assertTrue(out_docx.is_file())
        self.assertEqual(res["stats"]["headings"], 3)
        self.assertEqual(res["stats"]["tables"], 1)

        # Inspect resulting docx
        doc = docx.Document(str(out_docx))
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.tables[0].rows), 3)
        self.assertEqual(len(doc.tables[0].columns), 3)

    def test_callout_alerts_rendering(self):
        """Verify GitHub-style callouts (> [!NOTE], > [!WARNING]) render properly."""
        md_file = self.td / "callouts.md"
        content = """# Document with Callouts

> [!NOTE]
> This is a critical technical note regarding 5G Standalone operation.

> [!WARNING] Danger Ahead
> Failure to verify security credentials can lead to registration rejection.
"""
        md_file.write_text(content, encoding="utf-8")
        out_docx = self.td / "callouts.docx"

        convert_markdown_to_docx(md_path=md_file, output_docx_path=out_docx, apply_3gpp_styles=False)
        self.assertTrue(out_docx.is_file())

        doc = docx.Document(str(out_docx))
        # Callouts should be present as indented paragraphs with left borders
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any("NOTE" in t for t in texts))
        self.assertTrue(any("Danger Ahead" in t for t in texts))

    def test_embedded_media_and_companion_folder(self):
        """Verify figures referenced in Markdown are resolved and embedded in docx."""
        md_file = self.td / "spec_with_images.md"
        comp_media = self.td / "spec_with_images_media"
        comp_media.mkdir(parents=True, exist_ok=True)

        # Create a sample test image
        from PyQt5.QtGui import QImage, QColor
        img = QImage(200, 150, QImage.Format_ARGB32)
        img.fill(QColor(14, 116, 144)) # Cyan/teal
        img_path = comp_media / "arch_diagram.png"
        img.save(str(img_path), "PNG")

        content = """# Specification with Images

Below is the network architecture diagram:

![Figure 4.1: Reference Architecture](arch_diagram.png)
"""
        md_file.write_text(content, encoding="utf-8")
        out_docx = self.td / "spec_with_images.docx"

        res = convert_markdown_to_docx(
            md_path=md_file,
            output_docx_path=out_docx,
            companion_media_dir=comp_media,
        )

        self.assertEqual(res["stats"]["images"], 1)

        # Verify media file is stored in zip package
        with zipfile.ZipFile(str(out_docx), "r") as z:
            media_files = [f for f in z.namelist() if f.startswith("word/media/")]
            self.assertEqual(len(media_files), 1)

    def test_plantuml_diagram_compilation(self):
        """Verify embedded ```plantuml code blocks compile and embed as figures."""
        md_file = self.td / "diagram_test.md"
        content = """# Procedure with PlantUML

```plantuml
@startuml
skinparam BoxPadding 10
participant "UE" as UE
participant "AMF" as AMF
UE -> AMF: Registration Request
AMF -> UE: Registration Accept
@enduml
```
"""
        md_file.write_text(content, encoding="utf-8")
        out_docx = self.td / "diagram_test.docx"

        res = convert_markdown_to_docx(
            md_path=md_file,
            output_docx_path=out_docx,
            render_diagrams=True,
        )

        # If PlantUML and Java are available, diagram count should be 1
        if res["stats"]["diagrams"] > 0:
            with zipfile.ZipFile(str(out_docx), "r") as z:
                media_files = [f for f in z.namelist() if f.startswith("word/media/")]
                self.assertGreaterEqual(len(media_files), 1)

    def test_svg_to_png_conversion(self):
        """Verify SVG images are converted on the fly to PNG via QtSvg."""
        svg_file = self.td / "circle.svg"
        svg_file.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="120" height="120">'
            '<circle cx="60" cy="60" r="50" fill="#3B82F6"/>'
            '</svg>',
            encoding="utf-8",
        )
        png_out = self.td / "circle.png"

        rendered = render_svg_to_png(svg_file, png_out)
        self.assertTrue(rendered)
        self.assertTrue(png_out.is_file())
        self.assertGreater(png_out.stat().st_size, 0)

    def test_studio_ui_mode_switching_and_routing(self):
        """Verify WordMarkdownCard dual-mode switching and smart file drop routing."""
        card = WordMarkdownCard()

        # Defaults to mode 0 (Word → Markdown)
        self.assertEqual(card.studio_stack.currentIndex(), 0)
        self.assertTrue(card.btn_mode_w2m.isChecked())
        self.assertFalse(card.btn_mode_m2w.isChecked())

        # Switch to mode 1 (Markdown → Word)
        card.set_studio_mode(1)
        self.assertEqual(card.studio_stack.currentIndex(), 1)
        self.assertFalse(card.btn_mode_w2m.isChecked())
        self.assertTrue(card.btn_mode_m2w.isChecked())

        # Test toggle direction
        card._toggle_direction()
        self.assertEqual(card.studio_stack.currentIndex(), 0)

        # Create real temp files so DocumentSelectorPane sets ready state
        md_temp = self.td / "notes.md"
        md_temp.write_text("# Notes", encoding="utf-8")
        docx_temp = self.td / "draft.docx"
        docx.Document().save(str(docx_temp))

        # Test smart drop routing: dropping .md in w2m switches to mode 1 and clears w2m
        card._on_w2m_file_selected([str(md_temp)])
        self.assertEqual(card.studio_stack.currentIndex(), 1)
        self.assertEqual(card.pane_markdown.selected_files, [])
        self.assertEqual(card.pane_m2w.selected_files, [str(md_temp)])

        # Dropping .docx in m2w switches to mode 0 and clears m2w
        card._on_m2w_file_selected([str(docx_temp)])
        self.assertEqual(card.studio_stack.currentIndex(), 0)
        self.assertEqual(card.pane_m2w.selected_files, [])
        self.assertEqual(card.pane_markdown.selected_files, [str(docx_temp)])

        # Unrelated .docx completion while in mode 0 should NOT hijack studio mode to 1
        card._awaiting_m2w = False
        card.on_conversion_success(str(docx_temp))
        self.assertEqual(card.studio_stack.currentIndex(), 0)

    def test_hyperlinks_nested_lists_and_inline_elements(self):
        """Verify clickable hyperlinks, nested lists, <br>, <u>, <sup>, <sub>, and <hr>."""
        md_file = self.td / "rich_elements.md"
        content = """# Rich Formatting Test

Visit [3GPP Portal](https://portal.3gpp.org) for specifications.
Line one<br>Line two with <u>underlined</u>, <ins>inserted</ins>, X<sup>2</sup> and H<sub>2</sub>O.

---

- Top item 1
  - Nested item 1a
  - Nested item 1b
- Top item 2
"""
        md_file.write_text(content, encoding="utf-8")
        out_docx = self.td / "rich_elements.docx"

        convert_markdown_to_docx(md_path=md_file, output_docx_path=out_docx, apply_3gpp_styles=False)
        doc = docx.Document(str(out_docx))

        # Check clickable hyperlink XML presence
        doc_xml = doc.element.xml
        self.assertIn("w:hyperlink", doc_xml)

        # Verify nested list items are separate paragraphs (not concatenated into parent item)
        bullet_texts = [p.text.strip() for p in doc.paragraphs if "item" in p.text]
        self.assertIn("Top item 1", bullet_texts)
        self.assertIn("Nested item 1a", bullet_texts)
        self.assertIn("Nested item 1b", bullet_texts)
        self.assertIn("Top item 2", bullet_texts)

    def test_auto_detect_media_toggle(self):
        """Verify auto_detect_media controls companion folder lookup."""
        md_file = self.td / "doc.md"
        comp_dir = self.td / "doc_media"
        comp_dir.mkdir(parents=True, exist_ok=True)
        img_file = comp_dir / "fig1.png"
        img_file.write_bytes(b"dummy")

        # When auto_detect_media=True, finds fig1.png inside doc_media/
        resolved_on = resolve_image_path("fig1.png", md_file, auto_detect_media=True)
        self.assertIsNotNone(resolved_on)
        self.assertEqual(resolved_on.resolve(), img_file.resolve())

        # When auto_detect_media=False, does not search companion folder
        resolved_off = resolve_image_path("fig1.png", md_file, auto_detect_media=False)
        self.assertIsNone(resolved_off)

    def test_word_extractor_tab_integration(self):
        """Verify WordExtractorTab provides bidirectional Markdown studio integration."""
        tab = WordExtractorTab()

        # Ribbon nav button text and tooltip
        self.assertIn("Markdown ⇄ Word", tab.nav_btn_markdown.text())
        self.assertTrue(hasattr(tab, "convert_markdown_to_word_requested"))

        # Verify card aliases exist
        self.assertTrue(hasattr(tab, "pane_markdown"))
        self.assertTrue(hasattr(tab, "pane_m2w"))
        self.assertTrue(hasattr(tab, "run_m2w_btn"))


if __name__ == "__main__":
    unittest.main()
