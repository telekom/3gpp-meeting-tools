# --- File: tests/test_excel_company_review_export.py ---
import os
import sys
import unittest
import tempfile
from pathlib import Path

import openpyxl
from PyQt5.QtWidgets import QApplication

# Set offscreen QPA for headless testing
os.environ["QT_QPA_PLATFORM"] = "offscreen"
app = QApplication.instance() or QApplication(sys.argv)

from modules.meetings.core.excel_exporter import ExcelExporterThread
from modules.meetings.ui.excel_export_dialog import (
    ExcelExportDialog,
    ALL_AVAILABLE_COLUMNS,
    DEFAULT_SELECTED_COLUMNS,
    REVIEW_MATRIX_COLUMNS,
)


class TestExcelCompanyReviewExport(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.export_path = Path(self.tmp_dir.name) / "test_review_export.xlsx"

        self.sample_rows = [
            {
                "TDoc": "SP-260101",
                "Title": "Discussion on 5G Advanced Protocol Extensions",
                "Source": "Deutsche Telekom",
                "Type": "CR",
                "Agenda Item": "5.1",
                "TDoc Status": "Agreed",
                "For": "Agreement",
                "Secretary Remarks": "None",
                "My Status": "🟢 Support",
                "My Notes": "Key priority for our core network deployment.",
            },
            {
                "TDoc": "SP-260102",
                "Title": "Proposed changes to architectural assumptions",
                "Source": "Vendor A",
                "Type": "CR",
                "Agenda Item": "5.2",
                "TDoc Status": "Noted",
                "For": "Discussion",
                "Secretary Remarks": "Revised to SP-260199",
                "My Status": "🔴 Object",
                "My Notes": "Introduces backward-incompatible IE.",
            },
            {
                "TDoc": "SP-260103",
                "Title": "WID on 6G Phase 1 Feasibility",
                "Source": "Operator Alliance",
                "Type": "WID",
                "Agenda Item": "6.0",
                "TDoc Status": "Open",
                "For": "Approval",
                "Secretary Remarks": "",
                "My Status": "🟡 Monitor",
                "My Notes": "Wait for RAN1/RAN2 cross-working group feedback.",
            },
            {
                "TDoc": "SP-260104",
                "Title": "Correction on UPF Selection",
                "Source": "Deutsche Telekom, Orange",
                "Type": "CR",
                "Agenda Item": "5.3",
                "TDoc Status": "Open",
                "For": "Agreement",
                "Secretary Remarks": "",
                "My Status": "🔵 My TDoc",
                "My Notes": "Our co-sponsored contribution.",
            },
            {
                "TDoc": "SP-260105",
                "Title": "Report of Vice Chairman",
                "Source": "Vice Chairman",
                "Type": "Report",
                "Agenda Item": "2.1",
                "TDoc Status": "Noted",
                "For": "Information",
                "Secretary Remarks": "",
                "My Status": "⚪ Neutral",
                "My Notes": "",
            },
        ]

    def tearDown(self):
        self.tmp_dir.cleanup()

    def test_excel_export_dialog_presets_and_headers(self):
        """Tests ExcelExportDialog review matrix preset and company header mappings."""
        dialog = ExcelExportDialog(None, visible_count=5, total_count=10)

        # 1. Verify default company attribution
        self.assertTrue(dialog.is_company_attribution_enabled())
        dialog.edit_company.setText("Deutsche Telekom")
        headers = dialog.get_column_headers()
        self.assertEqual(headers["My Status"], "Status (Deutsche Telekom)")
        self.assertEqual(headers["My Notes"], "Comments (Deutsche Telekom)")

        # 2. Test Review Matrix preset
        dialog._apply_review_preset()
        selected = dialog.get_selected_columns()
        self.assertIn("My Status", selected)
        self.assertIn("My Notes", selected)
        self.assertEqual(selected, REVIEW_MATRIX_COLUMNS)

        # 3. Test Standard Agenda preset
        dialog._apply_standard_preset()
        selected_std = dialog.get_selected_columns()
        self.assertEqual(selected_std, DEFAULT_SELECTED_COLUMNS)
        self.assertNotIn("My Status", selected_std)
        self.assertNotIn("My Notes", selected_std)

        # 4. Test custom company name attribution
        dialog.edit_company.setText("Qualcomm")
        headers_qc = dialog.get_column_headers()
        self.assertEqual(headers_qc["My Status"], "Status (Qualcomm)")
        self.assertEqual(headers_qc["My Notes"], "Comments (Qualcomm)")

        # 5. Test toggling attribution off
        dialog.chk_company_attribution.setChecked(False)
        headers_off = dialog.get_column_headers()
        self.assertEqual(headers_off["My Status"], "My Status")
        self.assertEqual(headers_off["My Notes"], "My Notes")

    def test_excel_exporter_review_matrix_generation(self):
        """Tests ExcelExporterThread generation with company headers and status cell styling."""
        selected_columns = [
            "TDoc",
            "Title",
            "Source",
            "Agenda Item",
            "TDoc Status",
            "My Status",
            "My Notes",
        ]
        column_headers = {
            "My Status": "Status (Deutsche Telekom)",
            "My Notes": "Comments (Deutsche Telekom)",
        }

        exporter = ExcelExporterThread(
            output_path=self.export_path,
            rows_data=self.sample_rows,
            selected_columns=selected_columns,
            mtg_info={"wg_name": "SA2", "meeting_number": "168"},
            docs_ftp_url="https://www.3gpp.org/ftp/tsg_sa/WG2_Arch/TSGS2_168_Jeju/Docs",
            auto_open=False,
            column_headers=column_headers,
            company_name="Deutsche Telekom",
        )

        # Run synchronously
        exporter.run()

        self.assertTrue(self.export_path.exists())

        wb = openpyxl.load_workbook(str(self.export_path))
        ws = wb.active

        # Check Sheet title
        self.assertEqual(ws.title, "SA2 168")

        # 1. Verify Header row
        header_values = [ws.cell(row=1, column=col).value for col in range(1, len(selected_columns) + 1)]
        expected_headers = [
            "TDoc",
            "Title",
            "Source",
            "Agenda Item",
            "TDoc Status",
            "Status (Deutsche Telekom)",
            "Comments (Deutsche Telekom)",
        ]
        self.assertEqual(header_values, expected_headers)

        # 2. Verify Status column styling & values (Column 6)
        status_col = 6
        notes_col = 7

        # Row 2: Support
        r2_status = ws.cell(row=2, column=status_col)
        self.assertEqual(r2_status.value, "🟢 Support")
        self.assertEqual(r2_status.fill.start_color.rgb, "00E6F4EA")
        self.assertEqual(r2_status.font.color.rgb, "00137333")
        self.assertTrue(r2_status.font.bold)
        self.assertEqual(ws.cell(row=2, column=notes_col).value, "Key priority for our core network deployment.")

        # Row 3: Object
        r3_status = ws.cell(row=3, column=status_col)
        self.assertEqual(r3_status.value, "🔴 Object")
        self.assertEqual(r3_status.fill.start_color.rgb, "00FCE8E6")
        self.assertEqual(r3_status.font.color.rgb, "00C5221F")
        self.assertTrue(r3_status.font.bold)

        # Row 4: Monitor
        r4_status = ws.cell(row=4, column=status_col)
        self.assertEqual(r4_status.value, "🟡 Monitor")
        self.assertEqual(r4_status.fill.start_color.rgb, "00FEF7E0")
        self.assertEqual(r4_status.font.color.rgb, "00B06000")
        self.assertTrue(r4_status.font.bold)

        # Row 5: My TDoc
        r5_status = ws.cell(row=5, column=status_col)
        self.assertEqual(r5_status.value, "🔵 My TDoc")
        self.assertEqual(r5_status.fill.start_color.rgb, "00E8F0FE")
        self.assertEqual(r5_status.font.color.rgb, "001A73E8")
        self.assertTrue(r5_status.font.bold)

        # Row 6: Neutral -> should be empty string or None
        r6_status = ws.cell(row=6, column=status_col)
        self.assertIn(r6_status.value, ["", None])

        # 3. Verify column dimensions
        status_letter = openpyxl.utils.get_column_letter(status_col)
        notes_letter = openpyxl.utils.get_column_letter(notes_col)
        self.assertEqual(ws.column_dimensions[status_letter].width, 18)
        self.assertEqual(ws.column_dimensions[notes_letter].width, 42)


if __name__ == "__main__":
    unittest.main()
