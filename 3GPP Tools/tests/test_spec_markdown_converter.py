import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx

from core.utils.paths import to_extended_path
from modules.specifications.core.spec_markdown_converter import (
    convert_spec_to_markdown,
    discover_spec_parts,
    extract_spec_archive,
    get_part_sort_key,
    build_spec_metadata_header,
)


class TestSpecMarkdownConverter(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = Path(tempfile.mkdtemp(prefix="3gpp_spec_test_"))
        self.spec_dir = self.tmp_dir / "23501"
        self.spec_dir.mkdir(parents=True, exist_ok=True)

    def tearDown(self):
        shutil.rmtree(to_extended_path(self.tmp_dir), ignore_errors=True)

    def _create_sample_docx(self, doc_path: Path, headings: list):
        doc = docx.Document()
        doc.add_heading("3GPP Technical Specification", level=0)
        doc.add_paragraph("Front matter paragraph text.")
        for h_text, h_level, body_text in headings:
            doc.add_heading(h_text, level=h_level)
            doc.add_paragraph(body_text)
        doc.save(str(doc_path))

    def test_single_spec_unified_mode(self):
        stem = "23501-h40"
        docx_path = self.spec_dir / f"{stem}.docx"
        self._create_sample_docx(
            docx_path,
            [
                ("1 Scope", 1, "Scope of the 5G system specification."),
                ("2 References", 1, "References section."),
                ("4 Architecture model and concepts", 1, "Overview of 5G core architecture."),
                ("4.1 General", 2, "General subclause concepts."),
            ],
        )

        meta = {
            "spec": "23.501",
            "version": "17.4.0",
            "release": "Rel-17",
            "title": "System architecture for the 5G System (5GS)",
            "wg": "SP-SA2",
        }

        result = convert_spec_to_markdown(self.spec_dir, stem, mode="single", meta=meta)
        self.assertTrue(result.exists())
        self.assertEqual(result.name, f"{stem}.md")

        content = result.read_text(encoding="utf-8")
        self.assertIn("3GPP TS 23.501", content)
        self.assertIn("v17.4.0", content)
        self.assertIn("Scope of the 5G system specification", content)
        self.assertIn("Overview of 5G core architecture", content)

    def test_single_spec_clause_split_mode(self):
        stem = "23501-h40"
        docx_path = self.spec_dir / f"{stem}.docx"
        self._create_sample_docx(
            docx_path,
            [
                ("1 Scope", 1, "Scope of the 5G system specification."),
                ("2 References", 1, "References section."),
                ("4 Architecture model and concepts", 1, "Overview of 5G core architecture."),
                ("Annex A (normative): Parameters", 1, "Annex A content."),
            ],
        )

        meta = {
            "spec": "23.501",
            "version": "17.4.0",
            "release": "Rel-17",
            "title": "System architecture for the 5G System (5GS)",
        }

        index_result = convert_spec_to_markdown(self.spec_dir, stem, mode="split", meta=meta)
        self.assertTrue(index_result.exists())
        self.assertEqual(index_result.name, "index.md")

        clauses_dir = self.spec_dir / f"{stem}_clauses"
        self.assertTrue(clauses_dir.is_dir())

        fm_file = clauses_dir / "00_Front_Matter.md"
        c1_file = clauses_dir / "01_Scope.md"
        c4_file = clauses_dir / "04_Architecture_model_and_concepts.md"

        self.assertTrue(fm_file.exists(), "Front matter file should exist")
        self.assertTrue(c1_file.exists(), "Clause 1 file should exist")
        self.assertTrue(c4_file.exists(), "Clause 4 file should exist")

        index_content = index_result.read_text(encoding="utf-8")
        self.assertIn("[01_Scope.md](01_Scope.md)", index_content)
        self.assertIn("[04_Architecture_model_and_concepts.md](04_Architecture_model_and_concepts.md)", index_content)

    def test_multipart_spec_sequencing_and_stitching(self):
        stem = "38331-h40"
        p0_path = self.spec_dir / f"{stem}_s00_s04.docx"
        p1_path = self.spec_dir / f"{stem}_s05_s08.docx"

        self._create_sample_docx(
            p0_path,
            [
                ("1 Scope", 1, "NR RRC Scope."),
                ("4 General", 1, "General RRC model."),
            ],
        )
        self._create_sample_docx(
            p1_path,
            [
                ("5 Procedures", 1, "RRC connection establishment procedures."),
                ("6 Protocol data units", 1, "ASN.1 structure definitions."),
            ],
        )

        # Verify discovery & sort keys
        parts = discover_spec_parts(self.spec_dir, stem)
        self.assertEqual(len(parts), 2)
        self.assertEqual(parts[0].name, f"{stem}_s00_s04.docx")
        self.assertEqual(parts[1].name, f"{stem}_s05_s08.docx")

        meta = {"spec": "38.331", "version": "17.4.0", "release": "Rel-17"}

        # Test Mode 1 (Unified)
        unified_file = convert_spec_to_markdown(self.spec_dir, stem, mode="single", meta=meta)
        content = unified_file.read_text(encoding="utf-8")
        self.assertIn("NR RRC Scope", content)
        self.assertIn("RRC connection establishment procedures", content)

        # Test Mode 2 (Split across parts)
        index_file = convert_spec_to_markdown(self.spec_dir, stem, mode="split", meta=meta)
        clauses_dir = self.spec_dir / f"{stem}_clauses"
        self.assertTrue((clauses_dir / "01_Scope.md").exists())
        self.assertTrue((clauses_dir / "05_Procedures.md").exists())

    def test_media_isolation_across_versions(self):
        # Create two versions: Rel-17 (h40) and Rel-18 (i20)
        v1_stem = "23501-h40"
        v2_stem = "23501-i20"

        doc1 = self.spec_dir / f"{v1_stem}.docx"
        doc2 = self.spec_dir / f"{v2_stem}.docx"
        self._create_sample_docx(doc1, [("1 Scope", 1, "Version 17 text")])
        self._create_sample_docx(doc2, [("1 Scope", 1, "Version 18 text")])

        # Convert both in single mode
        res1 = convert_spec_to_markdown(self.spec_dir, v1_stem, mode="single")
        res2 = convert_spec_to_markdown(self.spec_dir, v2_stem, mode="single")

        # Media folders must be strictly separated
        v1_media = self.spec_dir / f"{v1_stem}_media"
        v2_media = self.spec_dir / f"{v2_stem}_media"

        self.assertTrue(v1_media.is_dir())
        self.assertTrue(v2_media.is_dir())
        self.assertNotEqual(v1_media, v2_media)

        # Convert both in split mode
        split_res1 = convert_spec_to_markdown(self.spec_dir, v1_stem, mode="split")
        split_res2 = convert_spec_to_markdown(self.spec_dir, v2_stem, mode="split")

        v1_clauses_media = self.spec_dir / f"{v1_stem}_clauses" / "media"
        v2_clauses_media = self.spec_dir / f"{v2_stem}_clauses" / "media"

        self.assertTrue(v1_clauses_media.is_dir())
        self.assertTrue(v2_clauses_media.is_dir())
        self.assertNotEqual(v1_clauses_media, v2_clauses_media)


if __name__ == "__main__":
    unittest.main()
