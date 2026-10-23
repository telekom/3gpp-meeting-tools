"""
Core utilities and background worker for extracting individual specification clauses.
Supports Word (.docx) excerpt generation with formatting/tables, Markdown export,
and multi-format clipboard copying (Rich HTML for Email/Outlook, Markdown, Plain Text).
"""

import html
import logging
import os
import shutil
import tempfile
from pathlib import Path
import re
from typing import Any, Dict, List, Optional

from PyQt5.QtCore import QMimeData, QThread, pyqtSignal
from PyQt5.QtWidgets import QApplication

from core.utils.paths import get_project_root, get_specs_root, to_extended_path
from modules.spec_search.core.spec_doc_resolver import (
    SpecDocResolver,
    match_multipart_document_for_clause,
)
from modules.word_tools.core.word_excerpt_extractor import (
    WordHeadingScanner,
    WordExcerptEngine,
)
from modules.meetings.core.tdoc_markdown_converter import convert_document_to_markdown

logger = logging.getLogger(__name__)


def get_default_clause_export_path(
    spec_number: str,
    version: str,
    clause_number: str,
    ext: str = ".docx",
    base_dir: Optional[Path] = None,
) -> Path:
    """Generates a clean, normalized destination file path for a clause export."""
    clean_spec = re.sub(r"^(?:TS|TR)\s*", "", str(spec_number).strip(), flags=re.IGNORECASE).strip()
    clean_ver = str(version).lstrip("vV").strip()
    clean_clause = str(clause_number).strip().replace(".", "_")

    if not base_dir:
        base_dir = get_specs_root() / clean_spec / "excerpts"

    os.makedirs(to_extended_path(base_dir), exist_ok=True)
    ext_clean = ext if ext.startswith(".") else f".{ext}"
    filename = f"TS_{clean_spec}_v{clean_ver}_Clause_{clean_clause}{ext_clean}"
    return base_dir / filename


def generate_clause_markdown(
    spec_number: str,
    version: str,
    clause_number: str,
    clause_title: str,
    content: str,
    release_date: Optional[str] = None,
) -> str:
    """Generates a standalone Markdown document representation of the clause."""
    clean_spec = str(spec_number).strip()
    clean_ver = str(version).lstrip("vV").strip()
    date_str = f" ({release_date})" if release_date else ""
    return (
        f"# 3GPP TS {clean_spec} v{clean_ver}{date_str}\n\n"
        f"## Clause {clause_number}: {clause_title}\n\n"
        f"{content.strip()}\n"
    )


def generate_clause_markdown_quote(
    spec_number: str,
    version: str,
    clause_number: str,
    clause_title: str,
    content: str,
    release_date: Optional[str] = None,
) -> str:
    """Generates a GitHub-flavored Markdown quote block suitable for notes and discussions."""
    clean_spec = str(spec_number).strip()
    clean_ver = str(version).lstrip("vV").strip()
    date_str = f" ({release_date})" if release_date else ""
    lines = content.strip().split("\n")
    quoted_lines = "\n".join(f"> {line}" for line in lines)
    return (
        f"> [!NOTE]\n"
        f"> **3GPP TS {clean_spec} v{clean_ver}{date_str} — Clause {clause_number}: {clause_title}**\n"
        f">\n"
        f"{quoted_lines}\n"
    )


def generate_clause_html_email(
    spec_number: str,
    version: str,
    clause_number: str,
    clause_title: str,
    content: str,
    release_date: Optional[str] = None,
) -> str:
    """
    Generates rich, styled HTML for Outlook, Microsoft Word, Gmail, and Teams.
    Includes a formal 3GPP citation header banner and properly spaced paragraphs.
    """
    clean_spec = str(spec_number).strip()
    clean_ver = str(version).lstrip("vV").strip()
    date_str = f" ({release_date})" if release_date else ""
    header_text = html.escape(f"3GPP TS {clean_spec} v{clean_ver}{date_str} — Clause {clause_number}: {clause_title}")

    # Format paragraphs preserving spacing and linebreaks
    paragraphs = content.strip().split("\n\n")
    body_parts = []
    for p in paragraphs:
        p_clean = p.strip()
        if not p_clean:
            continue
        p_escaped = html.escape(p_clean).replace("\n", "<br/>")
        body_parts.append(
            f'<p style="margin: 0 0 10px 0; font-family: Calibri, \'Segoe UI\', Arial, sans-serif; font-size: 11pt; line-height: 1.5; color: #1e293b;">{p_escaped}</p>'
        )
    body_html = "\n".join(body_parts)

    return f"""<div style="font-family: Calibri, 'Segoe UI', Arial, sans-serif; font-size: 11pt; color: #1e293b; max-width: 850px;">
  <div style="margin: 0 0 10px 0; padding: 8px 12px; background-color: #f1f5f9; border-left: 4px solid #0284c7; border-radius: 4px;">
    <strong style="color: #0369a1; font-size: 11.5pt;">{header_text}</strong>
  </div>
  <div style="padding-left: 4px;">
    {body_html}
  </div>
</div>"""


def copy_clause_to_clipboard(
    spec_number: str,
    version: str,
    clause_number: str,
    clause_title: str,
    content: str,
    release_date: Optional[str] = None,
    mode: str = "email",
) -> str:
    """
    Copies the clause to the clipboard according to the requested format.
    Modes:
      - 'email': Dual-format Rich HTML (for Outlook/Word) + Plain Text fallback
      - 'markdown': Markdown quote block with citation
      - 'text': Plain text with citation header
      - 'citation': Citation header only
    Returns a short user-facing status description (e.g. 'Copied for Email!').
    """
    clean_spec = str(spec_number).strip()
    clean_ver = str(version).lstrip("vV").strip()
    date_str = f" ({release_date})" if release_date else ""
    plain_citation = f"3GPP TS {clean_spec} v{clean_ver}{date_str}, Clause {clause_number} ({clause_title})"
    plain_full = f"{plain_citation}:\n\n{content.strip()}"

    clipboard = QApplication.clipboard()
    if not clipboard:
        return "Clipboard unavailable"

    if mode == "email":
        html_content = generate_clause_html_email(
            clean_spec, clean_ver, clause_number, clause_title, content, release_date
        )
        mime = QMimeData()
        mime.setHtml(html_content)
        mime.setText(plain_full)
        clipboard.setMimeData(mime)
        return "Copied for Email / Outlook!"

    elif mode == "markdown":
        md_text = generate_clause_markdown_quote(
            clean_spec, clean_ver, clause_number, clause_title, content, release_date
        )
        clipboard.setText(md_text)
        return "Copied as Markdown!"

    elif mode == "citation":
        clipboard.setText(plain_citation)
        return "Copied Citation!"

    else:  # plain text
        clipboard.setText(plain_full)
        return "Copied Plain Text!"


def export_clause_to_markdown_file(
    spec_number: str,
    version: str,
    clause_number: str,
    clause_title: str,
    content: str,
    release_date: Optional[str] = None,
    output_path: Optional[Path] = None,
) -> Path:
    """Writes the clause as a Markdown document to disk and returns the resolved Path."""
    if not output_path:
        output_path = get_default_clause_export_path(
            spec_number, version, clause_number, ext=".md"
        )
    output_path = Path(output_path).resolve()
    os.makedirs(to_extended_path(output_path.parent), exist_ok=True)

    md_content = generate_clause_markdown(
        spec_number, version, clause_number, clause_title, content, release_date
    )
    with open(to_extended_path(output_path), "w", encoding="utf-8") as f:
        f.write(md_content)

    return output_path


def create_clause_docx_fallback(
    spec_number: str,
    version: str,
    clause_number: str,
    clause_title: str,
    content: str,
    output_path: Path,
    release_date: Optional[str] = None,
) -> Path:
    """Fallback docx generator using python-docx when native heading extraction is unavailable."""
    import docx
    from docx.shared import Pt, RGBColor

    doc = docx.Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    date_str = f" ({release_date})" if release_date else ""
    header_para = doc.add_paragraph()
    header_run = header_para.add_run(f"3GPP TS {spec_number} v{version}{date_str}")
    header_run.bold = True
    header_run.font.size = Pt(11)
    header_run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    h1 = doc.add_heading(f"{clause_number} {clause_title}", level=1)
    h1.style.font.size = Pt(13)
    h1.style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    paragraphs = content.strip().split("\n\n")
    for p_text in paragraphs:
        p_clean = p_text.strip()
        if p_clean:
            doc.add_paragraph(p_clean)

    os.makedirs(to_extended_path(output_path.parent), exist_ok=True)
    doc.save(str(to_extended_path(output_path)))
    return output_path


class SpecClauseWordExtractorThread(QThread):
    """
    Background worker that extracts an individual clause from a 3GPP Word specification.
    Preserves native styles, formatting, and tables using WordExcerptEngine,
    falling back gracefully to structured python-docx generation if needed.
    """

    progress = pyqtSignal(int, str)
    finished_success = pyqtSignal(object)  # Path
    error = pyqtSignal(str)

    def __init__(
        self,
        spec_number: str,
        version: str,
        clause_number: str,
        clause_title: str,
        clause_content: str = "",
        output_path: Optional[Path] = None,
        cache_dir: Optional[Path] = None,
        specs_db: Optional[Any] = None,
        release_date: Optional[str] = None,
        export_format: str = "word",
    ):
        super().__init__()
        self.spec_number = str(spec_number).strip()
        self.version = str(version).lstrip("vV").strip()
        self.clause_number = str(clause_number).strip()
        self.clause_title = str(clause_title).strip()
        self.clause_content = clause_content.strip()
        self.output_path = output_path
        self.cache_dir = cache_dir
        self.specs_db = specs_db
        self.release_date = release_date
        self.export_format = export_format.lower().strip()
        self.images_count: int = 0

    def _find_matching_heading(self, headings: List[Any], clause_number: str) -> Optional[Any]:
        """Matches heading starting with the exact clause number followed by whitespace or tab."""
        clean_num = clause_number.strip()
        pattern = re.compile(r"^" + re.escape(clean_num) + r"(?:[\s\t]|$)", re.IGNORECASE)
        for h in headings:
            if pattern.match(h.text.strip()):
                return h
        return None

    def run(self):
        try:
            self.progress.emit(10, f"Locating specification document for TS {self.spec_number} v{self.version}...")

            if not self.output_path:
                ext = ".md" if self.export_format == "markdown" else ".docx"
                self.output_path = get_default_clause_export_path(
                    self.spec_number, self.version, self.clause_number, ext=ext
                )

            dest_path = Path(self.output_path).resolve()
            os.makedirs(to_extended_path(dest_path.parent), exist_ok=True)

            # 1. Resolve local files on disk
            info = SpecDocResolver.resolve_spec_files(
                self.spec_number, self.version, self.cache_dir, self.specs_db
            )
            spec_dir = info["spec_dir"]

            word_docs = info["word_docs"]
            # Extract from ZIP if only archive is present
            if not word_docs and info["zip_file"] and info["zip_file"].exists():
                self.progress.emit(25, f"Extracting archive {info['zip_file'].name}...")
                word_docs = SpecDocResolver.extract_word_from_zip(info["zip_file"], spec_dir)
            elif not word_docs and info.get("url"):
                from core.network.session import NetworkSession
                zip_target = spec_dir / (info.get("filename") or f"{info.get('stem', 'spec')}.zip")
                self.progress.emit(20, f"Downloading TS {self.spec_number} archive from 3GPP FTP...")
                try:
                    NetworkSession.download_file(info["url"], zip_target)
                    if zip_target.exists():
                        info["zip_file"] = zip_target
                        word_docs = SpecDocResolver.extract_word_from_zip(zip_target, spec_dir)
                except Exception as e:
                    logger.warning(f"Could not download ZIP archive for extraction: {e}")

            # 2. Extract via WordExcerptEngine if Word files are present
            if word_docs:
                target_doc = match_multipart_document_for_clause(word_docs, self.clause_number)
                source_paths = [str(target_doc)] if target_doc else [str(p) for p in word_docs]

                self.progress.emit(45, f"Scanning headings in {len(source_paths)} document(s)...")
                headings = WordHeadingScanner().scan(source_paths)
                target_h = self._find_matching_heading(headings, self.clause_number)

                if target_h:
                    if self.export_format == "markdown":
                        self.progress.emit(65, f"Extracting Clause {self.clause_number} from Word document...")
                        tmp_dir = Path(tempfile.mkdtemp(prefix="clause_md_export_"))
                        try:
                            temp_docx = tmp_dir / f"clause_{self.clause_number.replace('.', '_')}.docx"
                            WordExcerptEngine(source_paths, headings).extract([target_h.id], str(temp_docx))

                            self.progress.emit(85, "Converting clause to Markdown and extracting pictures...")
                            media_dir = dest_path.parent / f"{dest_path.stem}_media"
                            media_rel = f"{dest_path.stem}_media"

                            exported = convert_document_to_markdown(
                                temp_docx,
                                out_dir=dest_path.parent,
                                tdoc_id=dest_path.stem,
                                revisions="accept",
                                media_dir=media_dir,
                                media_rel=media_rel,
                            )

                            md_text = exported.text
                            header_end = md_text.find("\n---\n")
                            if header_end != -1 and md_text.startswith("#"):
                                body_text = md_text[header_end + 5:].lstrip()
                            else:
                                body_text = md_text

                            date_str = f" ({self.release_date})" if self.release_date else ""
                            final_md = (
                                f"# 3GPP TS {self.spec_number} v{self.version}{date_str}\n\n"
                                f"## Clause {self.clause_number}: {self.clause_title}\n\n"
                                f"{body_text.strip()}\n"
                            )

                            with open(to_extended_path(dest_path), "w", encoding="utf-8") as f:
                                f.write(final_md)

                            self.images_count = exported.images

                            if media_dir.exists() and (exported.images == 0 or not os.listdir(to_extended_path(media_dir))):
                                shutil.rmtree(to_extended_path(media_dir), ignore_errors=True)

                            self.progress.emit(100, f"Markdown excerpt ready ({self.images_count} image(s)).")
                            self.finished_success.emit(dest_path)
                            return
                        finally:
                            shutil.rmtree(to_extended_path(tmp_dir), ignore_errors=True)
                    else:
                        self.progress.emit(70, f"Extracting Clause {self.clause_number} via WordExcerptEngine...")
                        WordExcerptEngine(source_paths, headings).extract([target_h.id], str(dest_path))
                        self.images_count = 0
                        self.progress.emit(100, "Excerpt ready.")
                        self.finished_success.emit(dest_path)
                        return

            # 3. Fallback: Generate from clause content
            if self.clause_content:
                if self.export_format == "markdown":
                    self.progress.emit(80, "Generating Markdown excerpt from clause content...")
                    export_clause_to_markdown_file(
                        self.spec_number,
                        self.version,
                        self.clause_number,
                        self.clause_title,
                        self.clause_content,
                        release_date=self.release_date,
                        output_path=dest_path,
                    )
                    self.images_count = 0
                    self.progress.emit(100, "Markdown excerpt ready.")
                    self.finished_success.emit(dest_path)
                    return
                else:
                    self.progress.emit(80, "Generating clean Word excerpt from clause content...")
                    create_clause_docx_fallback(
                        self.spec_number,
                        self.version,
                        self.clause_number,
                        self.clause_title,
                        self.clause_content,
                        dest_path,
                        release_date=self.release_date,
                    )
                    self.images_count = 0
                    self.progress.emit(100, "Excerpt ready.")
                    self.finished_success.emit(dest_path)
                    return

            self.error.emit(
                f"Could not extract Clause {self.clause_number}: neither the Word document nor clause text was found."
            )

        except Exception as e:
            logger.error(f"Error extracting clause: {e}", exc_info=True)
            self.error.emit(str(e))


SpecClauseExtractorThread = SpecClauseWordExtractorThread
