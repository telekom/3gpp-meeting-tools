"""
Clause Word Document Comparator.
Coordinates clause extraction from Base and Target specification Word documents
via WordExcerptEngine, applies corporate sensitivity labels and 3GPP styles,
and automates Microsoft Word COM CompareDocuments to generate visual tracked changes.
"""

import logging
import os
import re
import shutil
import tempfile
from pathlib import Path
from typing import Any, List, Optional

import pythoncom
from PyQt5.QtCore import QThread, pyqtSignal

from modules.spec_search.core.spec_doc_resolver import (
    SpecDocResolver,
    match_multipart_document_for_clause,
)
from modules.word_tools.core.word_comparator import WordComparatorThread
from modules.word_tools.core.word_excerpt_extractor import (
    WordExcerptEngine,
    WordHeadingScanner,
)

logger = logging.getLogger(__name__)


class SpecClauseWordCompareThread(QThread):
    """
    Background worker that extracts a specific clause from two specification releases
    and generates a native Microsoft Word comparison document with tracked changes.
    """

    progress = pyqtSignal(int, str)
    finished_success = pyqtSignal(str)
    error = pyqtSignal(str)

    def __init__(
        self,
        spec_number: str,
        base_version: str,
        target_version: str,
        clause_number: str,
        clause_title: str = "",
        cache_dir: Optional[Path] = None,
        specs_db: Optional[Any] = None,
    ):
        super().__init__()
        self.spec_number = str(spec_number).strip()
        self.base_version = str(base_version).lstrip("v").strip()
        self.target_version = str(target_version).lstrip("v").strip()
        self.clause_number = str(clause_number).strip()
        self.clause_title = str(clause_title).strip()
        self.cache_dir = cache_dir
        self.specs_db = specs_db
        self._temp_dir: Optional[Path] = None

    def _find_matching_heading(self, headings: List[Any], clause_number: str) -> Optional[Any]:
        """Matches heading starting with the exact clause number followed by whitespace/tab."""
        clean_target = clause_number.strip()
        pattern = re.compile(r"^" + re.escape(clean_target) + r"(?:[\s\t]|$)", re.IGNORECASE)
        for h in headings:
            if pattern.match(h.text.strip()):
                return h
        return None

    def _create_placeholder_doc(self, path: Path, title: str, note: str):
        """Creates a minimal DOCX placeholder when a clause is absent in a release."""
        from docx import Document
        from modules.word_tools.core.sensitivity_label import set_sensitivity_label

        doc = Document()
        doc.add_heading(title, level=1)
        doc.add_paragraph(note)
        doc.save(str(path))

        try:
            import win32com.client
            pythoncom.CoInitialize()
            word = win32com.client.Dispatch("Word.Application")
            try:
                word_doc = word.Documents.Open(str(path))
                set_sensitivity_label(word_doc)
                word_doc.Save()
                word_doc.Close()
            finally:
                word = None
                pythoncom.CoUninitialize()
        except Exception as e:
            logger.debug(f"Placeholder label setting warning: {e}")

    def run(self):
        self._temp_dir = Path(tempfile.mkdtemp(prefix="clause_word_cmp_"))
        try:
            self.progress.emit(10, f"Locating documents for TS {self.spec_number}...")

            # 1. Resolve base and target files on disk
            base_info = SpecDocResolver.resolve_spec_files(
                self.spec_number, self.base_version, self.cache_dir, self.specs_db
            )
            target_info = SpecDocResolver.resolve_spec_files(
                self.spec_number, self.target_version, self.cache_dir, self.specs_db
            )

            # Extract from ZIP if only archive is present
            base_docs = base_info["word_docs"]
            if not base_docs and base_info["zip_file"] and base_info["zip_file"].exists():
                self.progress.emit(15, f"Extracting Base archive v{self.base_version}...")
                base_docs = SpecDocResolver.extract_word_from_zip(
                    base_info["zip_file"], base_info["spec_dir"]
                )

            target_docs = target_info["word_docs"]
            if not target_docs and target_info["zip_file"] and target_info["zip_file"].exists():
                self.progress.emit(25, f"Extracting Target archive v{self.target_version}...")
                target_docs = SpecDocResolver.extract_word_from_zip(
                    target_info["zip_file"], target_info["spec_dir"]
                )

            if not base_docs:
                self.error.emit(
                    f"Could not locate Word document for Base version v{self.base_version}.\n"
                    f"Please ensure it is downloaded in the Specifications tab."
                )
                return

            if not target_docs:
                self.error.emit(
                    f"Could not locate Word document for Target version v{self.target_version}.\n"
                    f"Please ensure it is downloaded in the Specifications tab."
                )
                return

            # 2. Extract Base Clause
            self.progress.emit(35, f"Scanning headings in Base release v{self.base_version}...")
            base_doc_target = match_multipart_document_for_clause(base_docs, self.clause_number)
            base_paths = [str(base_doc_target)] if base_doc_target else [str(p) for p in base_docs]

            base_headings = WordHeadingScanner().scan(base_paths)
            base_target_h = self._find_matching_heading(base_headings, self.clause_number)

            base_excerpt_path = self._temp_dir / f"Base_v{self.base_version}_Clause_{self.clause_number.replace('.', '_')}.docx"

            if base_target_h:
                self.progress.emit(50, f"Extracting Clause {self.clause_number} from Base v{self.base_version}...")
                WordExcerptEngine(base_paths, base_headings).extract([base_target_h.id], str(base_excerpt_path))
            else:
                logger.info(f"Clause {self.clause_number} was not present in Base release v{self.base_version}")
                self._create_placeholder_doc(
                    base_excerpt_path,
                    f"Clause {self.clause_number} {self.clause_title}",
                    f"[This clause was not present in 3GPP TS {self.spec_number} v{self.base_version}]",
                )

            # 3. Extract Target Clause
            self.progress.emit(65, f"Scanning headings in Target release v{self.target_version}...")
            target_doc_target = match_multipart_document_for_clause(target_docs, self.clause_number)
            target_paths = [str(target_doc_target)] if target_doc_target else [str(p) for p in target_docs]

            target_headings = WordHeadingScanner().scan(target_paths)
            target_target_h = self._find_matching_heading(target_headings, self.clause_number)

            target_excerpt_path = self._temp_dir / f"Target_v{self.target_version}_Clause_{self.clause_number.replace('.', '_')}.docx"

            if target_target_h:
                self.progress.emit(80, f"Extracting Clause {self.clause_number} from Target v{self.target_version}...")
                WordExcerptEngine(target_paths, target_headings).extract([target_target_h.id], str(target_excerpt_path))
            else:
                logger.info(f"Clause {self.clause_number} was removed in Target release v{self.target_version}")
                self._create_placeholder_doc(
                    target_excerpt_path,
                    f"Clause {self.clause_number} {self.clause_title}",
                    f"[This clause was removed in 3GPP TS {self.spec_number} v{self.target_version}]",
                )

            # 4. Compare via Microsoft Word COM (Legal Compare with Track Changes)
            self.progress.emit(90, "Launching Microsoft Word Legal Compare (Track Changes)...")
            comparator = WordComparatorThread(
                str(base_excerpt_path),
                str(target_excerpt_path),
                keep_open=False,
            )
            comparator.run()

            self.progress.emit(100, "Comparison complete!")
            self.finished_success.emit(
                f"Generated Word comparison with Track Changes for Clause {self.clause_number}: v{self.base_version} vs v{self.target_version}"
            )
        except Exception as e:
            logger.error(f"Failed to generate clause comparison in Word: {e}", exc_info=True)
            self.error.emit(str(e))
        finally:
            # Note: We do not immediately wipe _temp_dir if Word needs to view the comparison,
            # or we leave cleanup to OS temp
            pass
