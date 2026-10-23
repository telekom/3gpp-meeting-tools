# --- File: src/modules/meetings/core/meeting_wi_enricher.py ---
import concurrent.futures
import json
import logging
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, Optional, Set

import openpyxl
from PyQt5.QtCore import QThread, pyqtSignal

from core.utils.paths import get_project_root
from modules.meetings.core.settings import MeetingsSettings
from modules.work_items.core.wi_database import WorkItemsDatabase
from modules.work_items.core.wi_scraper import fetch_wi_details_from_portal

logger = logging.getLogger(__name__)


def get_marker_path(excel_filepath: Path) -> Path:
    """Returns the path for the enrichment marker JSON file next to the Excel."""
    path = Path(excel_filepath)
    return path.parent / f"{path.name}.wi_checked.json"


def is_enrichment_needed(excel_filepath: Path) -> bool:
    """
    Returns True if the Excel has not been checked yet or if the Excel file
    was modified/re-downloaded after the last check.
    """
    path = Path(excel_filepath)
    if not path.exists():
        return False
    marker = get_marker_path(path)
    if not marker.exists():
        return True
    try:
        return os.path.getmtime(marker) < os.path.getmtime(path)
    except Exception:
        return True


def extract_wi_ids_from_excel(
    excel_filepath: Path, cancel_check: Optional[Callable[[], bool]] = None
) -> Dict[str, str]:
    """
    Scans the Excel file off-thread for Work Item hyperlinks in single-WI cells.
    Returns a dict mapping numeric WI codes (strings) to their display values/acronyms.
    """
    path = Path(excel_filepath)
    if not path.exists():
        return {}

    found: Dict[str, str] = {}
    wb = None
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
        sheet = wb.worksheets[0]

        header_row = 1
        wi_col = None
        tdoc_col = None

        # Header hunt (check first 15 rows)
        for r in range(1, min(16, sheet.max_row + 1)):
            for c in range(1, min(sheet.max_column + 1, 50)):
                val = str(sheet.cell(row=r, column=c).value or "").strip().upper()
                if not wi_col and ("RELATED WI" in val or "WORK ITEM" in val or val in ["WI", "WID"]):
                    wi_col = c
                if not tdoc_col and (val in ["TDOC", "TD#", "TDOC#"] or "TDOC" in val):
                    tdoc_col = c
            if wi_col:
                header_row = r
                break

        if not wi_col:
            logger.debug(f"[WI Enrichment] No WI column found in {path.name}")
            return {}

        empty_streak = 0
        for r in range(header_row + 1, sheet.max_row + 1):
            if cancel_check and cancel_check():
                break

            # Stop scanning when TDoc rows end (safely bounded)
            if tdoc_col:
                tdoc_val = sheet.cell(row=r, column=tdoc_col).value
                if not tdoc_val:
                    empty_streak += 1
                    if empty_streak >= 10:
                        break
                    continue
                else:
                    empty_streak = 0

            cell = sheet.cell(row=r, column=wi_col)
            hl = cell.hyperlink.target if cell.hyperlink else None
            if hl:
                m = re.search(r"workitemId=(\d+)", str(hl), re.IGNORECASE)
                if m:
                    wi_id = m.group(1)
                    val_str = str(cell.value or "").strip()
                    if wi_id not in found:
                        found[wi_id] = val_str

    except Exception as e:
        logger.warning(f"[WI Enrichment] Error scanning Excel {path.name} for WI hyperlinks: {e}")
    finally:
        if wb:
            try:
                wb.close()
            except Exception:
                pass

    return found


def write_enrichment_marker(
    excel_filepath: Path,
    unique_wis_found: int,
    new_wis_scraped: int,
    scraped_codes: Optional[list] = None,
):
    """Writes the audit marker file indicating that this Excel revision was checked."""
    try:
        path = Path(excel_filepath)
        marker = get_marker_path(path)
        payload = {
            "checked_at": datetime.now().isoformat(),
            "excel_mtime": os.path.getmtime(path) if path.exists() else 0,
            "unique_wis_found": unique_wis_found,
            "new_wis_scraped": new_wis_scraped,
            "scraped_codes": scraped_codes or [],
        }
        with open(marker, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2)
    except Exception as e:
        logger.debug(f"[WI Enrichment] Failed to write marker file: {e}")


class MeetingWIEnricherThread(QThread):
    """
    Dedicated background worker that:
    1. Fast-exits if the meeting Excel was already checked.
    2. Off-thread scans the WI column for 3GPP Portal hyperlinks.
    3. Fast-queries SQLite to discover which WIs are missing from 3gpp_data.db.
    4. Concurrently scrapes and upserts missing WIs via 3GPP Portal.
    5. Writes a marker file so subsequent opens take ~0 ms.
    """

    finished_enrichment = pyqtSignal(bool, int, str)

    def __init__(self, excel_filepath: Path, db_path: Optional[Path] = None, parent=None):
        super().__init__(parent)
        self.excel_path = Path(excel_filepath)
        self.db_path = db_path or (get_project_root() / "db" / "3gpp_data.db")
        self.settings = MeetingsSettings()

    def run(self):
        try:
            if self.isInterruptionRequested():
                self.finished_enrichment.emit(False, 0, "Cancelled")
                return

            # Check user setting toggle
            if not self.settings.get_auto_enrich_wis():
                self.finished_enrichment.emit(True, 0, "Auto-enrichment disabled in settings")
                return

            if not is_enrichment_needed(self.excel_path):
                self.finished_enrichment.emit(True, 0, "Excel already enriched and up to date")
                return

            # Step 1: Extract WI IDs off-thread
            found_wis = extract_wi_ids_from_excel(
                self.excel_path, cancel_check=self.isInterruptionRequested
            )

            if self.isInterruptionRequested():
                self.finished_enrichment.emit(False, 0, "Cancelled")
                return

            if not found_wis:
                write_enrichment_marker(self.excel_path, unique_wis_found=0, new_wis_scraped=0)
                self.finished_enrichment.emit(True, 0, "No WI hyperlinks found")
                return

            # Step 2: Query database for existing WI codes
            db = WorkItemsDatabase(self.db_path)
            existing_codes = db.get_existing_wi_codes(set(found_wis.keys()))

            missing_codes = [c for c in found_wis if c not in existing_codes]

            if not missing_codes:
                write_enrichment_marker(
                    self.excel_path,
                    unique_wis_found=len(found_wis),
                    new_wis_scraped=0,
                )
                logger.debug(
                    f"[WI Enrichment] Meeting '{self.excel_path.name}': All {len(found_wis)} WIs already in DB."
                )
                self.finished_enrichment.emit(True, 0, "All WIs already in database")
                return

            if self.isInterruptionRequested():
                self.finished_enrichment.emit(False, 0, "Cancelled")
                return

            # Step 3: Concurrently fetch missing WIs (polite worker pool)
            logger.info(
                f"ℹ️ [WI Enrichment] Meeting '{self.excel_path.name}': Discovered {len(missing_codes)} "
                f"new Work Item(s) missing from database. Scraping in background..."
            )

            scraped_count = 0
            successful_codes = []
            max_workers = min(4, len(missing_codes))

            with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                future_to_code = {
                    executor.submit(fetch_wi_details_from_portal, code): code
                    for code in missing_codes
                }

                for future in concurrent.futures.as_completed(future_to_code):
                    if self.isInterruptionRequested():
                        break
                    code = future_to_code[future]
                    try:
                        meta = future.result()
                        if meta and not meta.get("error"):
                            if db.upsert_single_work_item(meta):
                                scraped_count += 1
                                successful_codes.append(code)
                                logger.debug(
                                    f"[WI Enrichment] Upserted WI #{code} ({meta.get('acronym')})"
                                )
                    except Exception as e:
                        logger.debug(f"[WI Enrichment] Error fetching WI #{code}: {e}")

            if self.isInterruptionRequested():
                self.finished_enrichment.emit(False, scraped_count, "Cancelled")
                return

            # Step 4: Write marker file
            write_enrichment_marker(
                self.excel_path,
                unique_wis_found=len(found_wis),
                new_wis_scraped=scraped_count,
                scraped_codes=successful_codes,
            )

            if scraped_count > 0:
                logger.info(
                    f"✅ [WI Enrichment] Meeting '{self.excel_path.name}': Successfully scraped and "
                    f"indexed {scraped_count} new Work Item(s) in database."
                )
            else:
                logger.debug(
                    f"[WI Enrichment] Meeting '{self.excel_path.name}': Completed with no new WIs saved."
                )

            self.finished_enrichment.emit(
                True, scraped_count, f"Successfully enriched {scraped_count} Work Item(s)"
            )

        except Exception as e:
            logger.warning(
                f"[WI Enrichment] Unexpected error enriching meeting '{self.excel_path.name}': {e}",
                exc_info=True,
            )
            self.finished_enrichment.emit(False, 0, str(e))
